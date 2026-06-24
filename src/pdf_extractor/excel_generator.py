# -*- coding: utf-8 -*-
"""Excel 生成器：基于 openpyxl 生成 .xlsx 文件。"""

import logging
import os
import re
import shutil
from typing import List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
HEADER_FONT = Font(bold=True)


class ExcelGenerator:
    """封装 openpyxl，提供统一的 Excel 生成接口。"""

    def __init__(self, columns: List[str]):
        """初始化 Excel 生成器。

        Args:
            columns: Excel 列标题列表，包含"来源文件"列
        """
        self.columns = columns

    def generate(self, data: List[dict], output_path: str, sort_key: str = "编号") -> bool:
        """生成 Excel 文件。

        Args:
            data: 要写入的数据列表
            output_path: 输出文件绝对路径
            sort_key: 排序字段名称，默认按"编号"升序

        Returns:
            生成成功返回 True，失败返回 False
        """
        try:
            sorted_data = self._sort_data(data, sort_key)
            wb = Workbook()
            ws = wb.active
            ws.title = "PDF信息汇总"

            # 写入表头
            for col_idx, col_name in enumerate(self.columns, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT

            # 写入数据
            for row_idx, record in enumerate(sorted_data, start=2):
                for col_idx, col_name in enumerate(self.columns, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=record.get(col_name, ""))

            # 自动列宽
            for col_idx, col_name in enumerate(self.columns, start=1):
                max_len = len(str(col_name))
                for row in ws.iter_rows(
                    min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx
                ):
                    for cell in row:
                        if cell.value is not None:
                            max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

            out_dir = os.path.dirname(os.path.abspath(output_path))
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)
            wb.save(output_path)
            logger.info("Excel 已生成: %s", output_path)
            return True
        except PermissionError:
            logger.error("Excel 文件被占用，无法写入: %s", output_path)
            return False
        except OSError as e:
            if "No space left" in str(e) or getattr(e, "errno", None) == 28:
                logger.error("磁盘空间不足")
            else:
                logger.exception("生成 Excel 失败: %s", e)
            return False
        except Exception as e:
            logger.exception("生成 Excel 失败: %s", e)
            return False

    @staticmethod
    def _sort_data(data: List[dict], sort_key: str) -> List[dict]:
        """按指定字段升序排序（支持数字前缀）；字段不存在时保持原序。"""
        if not sort_key or not any(sort_key in rec for rec in data):
            return list(data)

        def sort_value(record: dict):
            val = str(record.get(sort_key, ""))
            num_match = re.search(r"\d+", val)
            if num_match:
                return (0, int(num_match.group()), val)
            return (1, 0, val)

        return sorted(data, key=sort_value)

    def generate_from_template(
        self, data: List[dict], template_path: str, output_path: str,
        sort_key: str = "编号"
    ) -> bool:
        """基于模板生成 Excel — 复制模板后填入数据，保留格式。"""
        try:
            out_dir = os.path.dirname(os.path.abspath(output_path))
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)
            shutil.copy2(template_path, output_path)
            wb = load_workbook(output_path)
            ws = wb.active

            # 找表头行
            header_row = 1
            for row_idx in range(1, min(ws.max_row + 1, 50)):
                vals = [str(ws.cell(row=row_idx, column=c).value or "")
                        for c in range(1, ws.max_column + 1)]
                matches = sum(1 for cn in self.columns
                              if any(cn in v or v in cn for v in vals if v))
                if matches >= max(1, len(self.columns) * 0.5):
                    header_row = row_idx
                    break

            data_start = header_row + 1
            col_map = {}
            for ci in range(1, ws.max_column + 1):
                h = str(ws.cell(row=header_row, column=ci).value or "").strip()
                for oc in self.columns:
                    if oc in h or h in oc:
                        col_map[oc] = ci
                        break
            next_col = ws.max_column + 1
            for oc in self.columns:
                if oc not in col_map:
                    col_map[oc] = next_col
                    next_col += 1

            sorted_data = self._sort_data(data, sort_key)
            for ri, rec in enumerate(sorted_data):
                row_idx = data_start + ri
                for cn, ci in col_map.items():
                    ws.cell(row=row_idx, column=ci).value = rec.get(cn, "")

            for ci in range(1, ws.max_column + 1):
                max_len = max(
                    (len(str(ws.cell(row=r, column=ci).value or ""))
                     for r in range(1, ws.max_row + 1)),
                    default=0)
                if max_len > 0:
                    ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 4, 50)

            wb.save(output_path)
            logger.info("Excel (模板) 已生成: %s", output_path)
            return True
        except Exception as e:
            logger.exception("模板生成 Excel 失败: %s", e)
            return False

    @staticmethod
    def default_filename() -> str:
        """生成默认文件名。"""
        from datetime import datetime

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"PDF信息汇总表_{ts}.xlsx"
